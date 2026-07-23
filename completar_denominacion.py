import openpyxl
from openpyxl.styles import PatternFill

# Cargar archivos
print("Cargando archivos...")
wb_principal = openpyxl.load_workbook('../Pipetas Generales SWISSMEDICAL 19.06.2026 - ACTUALIZADO_FINAL 2.xlsx')
wb_sap = openpyxl.load_workbook('sap.xlsx')
wb_sap2 = openpyxl.load_workbook('sap2.xlsx')

# Leer sap2 (equipos que SÍ están en SAP)
ws_sap2 = wb_sap2.active
seriales_en_sap2 = set()
print(f"\nLeyendo sap2.xlsx...")
for row in range(2, ws_sap2.max_row + 1):
    serie = ws_sap2.cell(row, 3).value  # Columna C: Serie
    if serie:
        seriales_en_sap2.add(str(serie).strip().upper())

print(f"  Total seriales en SAP2: {len(seriales_en_sap2)}")

# Leer sap.xlsx para crear mapeo Magnitud -> Denomin.tipo
ws_sap = wb_sap.active
mapeo_magnitud_denominacion = {}

print(f"\nCreando mapeo desde sap.xlsx...")
print("  Primero extraemos las magnitudes del archivo principal...")

# Necesitamos primero extraer las magnitudes únicas del archivo principal
magnitudes_principales = {}
for hoja_nombre in wb_principal.sheetnames:
    ws = wb_principal[hoja_nombre]
    
    # Buscar columna de Serie y Magnitud
    headers = {}
    for col in range(1, ws.max_column + 1):
        header = ws.cell(1, col).value
        if header:
            headers[str(header).strip().upper()] = col
    
    col_serie = headers.get('SERIE', headers.get('NÚMERO DE SERIE', headers.get('NUMERO DE SERIE')))
    col_magnitud = headers.get('MAGNITUD', headers.get('VOLUMEN'))
    
    if col_serie and col_magnitud:
        for row in range(2, ws.max_row + 1):
            serie = ws.cell(row, col_serie).value
            magnitud = ws.cell(row, col_magnitud).value
            if serie and magnitud:
                magnitudes_principales[str(serie).strip().upper()] = str(magnitud).strip()

# Ahora buscamos los equipos de sap.xlsx y extraemos su magnitud del principal
for row in range(2, ws_sap.max_row + 1):
    equipo = str(ws_sap.cell(row, 1).value).strip()  # Columna A: Equipo
    denominacion = ws_sap.cell(row, 2).value  # Columna B: Denomin.tipo
    
    if equipo and denominacion:
        # Extraer el serial del equipo (ej: PIP-KJ15147 -> KJ15147)
        if '-' in equipo:
            serial = equipo.split('-')[-1].strip().upper()
        else:
            serial = equipo.strip().upper()
        
        # Buscar la magnitud de este serial en el archivo principal
        if serial in magnitudes_principales:
            magnitud = magnitudes_principales[serial]
            # Crear mapeo Magnitud -> Denomin.tipo
            if magnitud not in mapeo_magnitud_denominacion:
                mapeo_magnitud_denominacion[magnitud] = denominacion

print(f"\n  Mapeo creado con {len(mapeo_magnitud_denominacion)} magnitudes:")
for mag, denom in sorted(mapeo_magnitud_denominacion.items()):
    print(f"    {mag} -> {denom}")

# Procesar todas las hojas del archivo principal
fill_completado = PatternFill(start_color="00FFFF", end_color="00FFFF", fill_type="solid")  # Cyan

total_completados = 0
reporte_por_hoja = {}

print(f"\n{'='*60}")
print("PROCESANDO ARCHIVO PRINCIPAL")
print(f"{'='*60}")

for hoja_nombre in wb_principal.sheetnames:
    ws = wb_principal[hoja_nombre]
    print(f"\nHoja: {hoja_nombre}")
    
    # Buscar columnas
    headers = {}
    for col in range(1, ws.max_column + 1):
        header = ws.cell(1, col).value
        if header:
            headers[str(header).strip().upper()] = col
    
    # Buscar columna Serie, Magnitud y Denomin.tipo
    col_serie = headers.get('SERIE', headers.get('NÚMERO DE SERIE', headers.get('NUMERO DE SERIE')))
    col_magnitud = headers.get('MAGNITUD', headers.get('VOLUMEN'))
    col_denominacion = None
    
    # Buscar columna Denomin.tipo o similar
    for key in headers.keys():
        if 'DENOMIN' in key or 'DENOMINACION' in key:
            col_denominacion = headers[key]
            break
    
    # Si no existe la columna Denomin.tipo, crearla
    if not col_denominacion:
        col_denominacion = ws.max_column + 1
        ws.cell(1, col_denominacion).value = 'Denomin.tipo'
        print(f"  ✓ Columna 'Denomin.tipo' creada en columna {col_denominacion}")
    
    if not col_serie or not col_magnitud:
        print(f"  ⚠ No se encontraron columnas necesarias (Serie y/o Magnitud)")
        continue
    
    completados_hoja = 0
    
    # Procesar cada fila
    for row in range(2, ws.max_row + 1):
        serie = ws.cell(row, col_serie).value
        magnitud = ws.cell(row, col_magnitud).value
        denominacion_actual = ws.cell(row, col_denominacion).value
        
        if not serie:
            continue
        
        serie_clean = str(serie).strip().upper()
        
        # Si la serie NO está en SAP2 y tiene magnitud
        if serie_clean not in seriales_en_sap2 and magnitud:
            magnitud_clean = str(magnitud).strip()
            
            # Si no tiene Denomin.tipo o está vacío
            if not denominacion_actual or str(denominacion_actual).strip() == '':
                # Buscar en el mapeo
                if magnitud_clean in mapeo_magnitud_denominacion:
                    nueva_denominacion = mapeo_magnitud_denominacion[magnitud_clean]
                    ws.cell(row, col_denominacion).value = nueva_denominacion
                    ws.cell(row, col_denominacion).fill = fill_completado
                    completados_hoja += 1
                    print(f"    Fila {row}: {serie_clean} | {magnitud_clean} -> {nueva_denominacion}")
    
    if completados_hoja > 0:
        print(f"  ✓ Completados en esta hoja: {completados_hoja}")
        reporte_por_hoja[hoja_nombre] = completados_hoja
        total_completados += completados_hoja
    else:
        print(f"  ○ Sin cambios necesarios")

# Guardar
print(f"\n{'='*60}")
print("RESUMEN")
print(f"{'='*60}")
print(f"Total Denomin.tipo completados: {total_completados}")
print(f"\nPor hoja:")
for hoja, cantidad in reporte_por_hoja.items():
    print(f"  {hoja}: {cantidad}")

print(f"\n{'='*60}")
print("Guardando archivo...")
wb_principal.save('../Pipetas Generales SWISSMEDICAL 19.06.2026 - ACTUALIZADO_FINAL 2.xlsx')
print("✓ Archivo guardado correctamente")
print("\nLas celdas completadas están marcadas en color CYAN")

# Crear reporte
reporte = f"""
REPORTE DE COMPLETADO DE DENOMINACIONES
========================================

Archivos utilizados:
- Principal: Pipetas Generales SWISSMEDICAL 19.06.2026 - ACTUALIZADO_FINAL 2.xlsx
- SAP (no en sistema): sap.xlsx
- SAP2 (en sistema): sap2.xlsx

Lógica aplicada:
- Se identificaron {len(seriales_en_sap2)} seriales que SÍ están en SAP2
- Se creó un mapeo de {len(mapeo_magnitud_denominacion)} magnitudes -> Denomin.tipo desde sap.xlsx
- Se completó el campo "Denomin.tipo" para seriales que NO están en SAP2

RESULTADOS:
Total de Denomin.tipo completados: {total_completados}

Detalle por hoja:
"""

for hoja, cantidad in reporte_por_hoja.items():
    reporte += f"  - {hoja}: {cantidad} completados\n"

reporte += f"""
MAPEO UTILIZADO (Magnitud -> Denomin.tipo):
"""
for mag, denom in sorted(mapeo_magnitud_denominacion.items()):
    reporte += f"  {mag} = {denom}\n"

with open('REPORTE_DENOMINACION.txt', 'w', encoding='utf-8') as f:
    f.write(reporte)

print(f"\n✓ Reporte guardado en: REPORTE_DENOMINACION.txt")
